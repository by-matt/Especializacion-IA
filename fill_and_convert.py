import os
import docx
import win32com.client

def fill_and_export():
    src_docx = r"C:\Users\VN\Downloads\1_Formulario_Ficha_de_Alineacion_Tecnica.docx"
    dest_dir = r"C:\Users\VN\Downloads\Especializacion IA"
    out_docx = os.path.join(dest_dir, "Ficha_Alineacion_ByronCalderon.docx")
    out_pdf = os.path.join(dest_dir, "Ficha_Alineacion_ByronCalderon.pdf")

    doc = docx.Document(src_docx)

    # Tabla 0: Datos generales
    t0 = doc.tables[0]
    t0.rows[0].cells[1].text = "Personalizado"
    t0.rows[1].cells[1].text = (
        "Ingeniero en Biotecnología. No provengo del desarrollo de software tradicional, "
        "pero utilizo intensivamente herramientas de IA para resolver problemas técnicos y "
        "automatizar tareas. Mi objetivo es integrar perfiles: aprovechar mi base en biotecnología "
        "y análisis de datos, y aprender a construir servicios backend e integraciones de APIs "
        "con IA que sean robustos, útiles y testeados en producción."
    )

    # Tabla 1: Diagnóstico técnico
    t1 = doc.tables[1]
    t1.rows[1].cells[1].text = "3"
    t1.rows[1].cells[2].text = (
        "Me apoyo bastante en herramientas de IA para generar y depurar código; entiendo variables, "
        "funciones y control de flujo, pero busco ganar autonomía en la estructuración de scripts más "
        "complejos y manejo sistemático de excepciones."
    )
    t1.rows[2].cells[1].text = "2"
    t1.rows[2].cells[2].text = (
        "Conozco los comandos elementales (commit, push, clone), pero me cuesta el trabajo con ramas, "
        "la resolución de conflictos y entender el estado del repositorio sin asistencia paso a paso."
    )
    t1.rows[3].cells[1].text = "3"
    t1.rows[3].cells[2].text = (
        "He realizado peticiones a servicios guiado por asistentes de IA y comprendo el formato JSON "
        "y las API keys; mi brecha está en dominar el ciclo HTTP, el control de rate limits y reintentos "
        "por código propio."
    )
    t1.rows[4].cells[1].text = "3"
    t1.rows[4].cells[2].text = (
        "Me desenvuelvo en la terminal para lo esencial (instalar paquetes, activar entornos virtuales); "
        "cuando surgen conflictos de dependencias o variables de entorno, dependo de la IA para diagnosticarlos."
    )

    # Tabla 2: Problema de negocio tentativo
    t2 = doc.tables[2]
    t2.rows[0].cells[0].text = (
        "Ayuda a profesionales, analistas e investigadores en biotecnología a consultar y validar "
        "condiciones de bioprocesos o protocolos de laboratorio. Resuelve la tarea de buscar y extraer "
        "de forma confiable parámetros clave (rangos de pH, temperatura, sustratos y tasas de rendimiento) "
        "y normas de bioseguridad, respondiendo a partir de un conjunto de documentos técnicos, fichas y "
        "literatura especializada indexada mediante RAG."
    )

    # Tabla 3: Progresión de los 3 Gates
    t3 = doc.tables[3]
    t3.rows[0].cells[0].text = (
        "En el Gate 1 construyo la base técnica para comunicarme con un modelo vía API de forma confiable, "
        "logrando que procese entradas y responda en JSON sin caerse ante errores de red. En el Gate 2 no "
        "descarto ese código: lo convierto en el motor interno para que ahora consulte mis propios documentos "
        "de biotecnología (RAG) y tome decisiones mediante herramientas y agentes. Finalmente, en el Gate 3, "
        "tomo esa solución y la empaqueto como un producto real: la expongo como una API web con FastAPI, "
        "la meto dentro de un contenedor Docker para que corra en cualquier máquina sin configurar nada extra, "
        "y le agrego matrices de pruebas y control de riesgos."
    )

    # Tabla 4: Criterios críticos no compensables
    t4 = doc.tables[4]
    t4.rows[0].cells[1].text = (
        "Manejo de errores de API y no exponer secretos: Que el programa capture caídas de conexión "
        "o límites de consulta (rate limits) sin cerrarse abruptamente, y que las credenciales/API keys "
        "jamás se suban al repositorio de Git (.env en .gitignore)."
    )
    t4.rows[1].cells[1].text = (
        "Cero alucinación y respaldo en fuentes RAG: Que el sistema no invente datos y responda "
        "exclusivamente respaldado en los documentos indexados, citando la fuente o declarando con "
        "honestidad que la información no está disponible si no la encuentra."
    )
    t4.rows[2].cells[1].text = (
        "Ejecución 100% reproducible en Docker: Que la aplicación completa se pueda clonar y levantar "
        "en cualquier máquina únicamente con Docker, funcionando en sus endpoints sin requerir instalaciones "
        "manuales previas ni fallar por configuraciones locales."
    )

    # Tabla 5: Tu punto de partida declarado (fortalezas y brechas)
    t5 = doc.tables[5]
    t5.rows[0].cells[0].text = (
        "Fortalezas: Pensamiento analítico y metodológico propio de la biotecnología, y gran soltura "
        "aprovechando herramientas de IA para estructurar ideas, resolver problemas y crear código funcional.\n\n"
        "Brechas: Poca experiencia en desarrollo de software formal, especialmente en el dominio autónomo "
        "de Git, arquitectura de microservicios con FastAPI y despliegue contenerizado en Docker."
    )

    # Guardar DOCX modificado
    doc.save(out_docx)
    print(f"DOCX guardado exitosamente en: {out_docx}")

    # Exportar a PDF usando Word COM
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        abs_in = os.path.abspath(out_docx)
        abs_out = os.path.abspath(out_pdf)
        doc_com = word.Documents.Open(abs_in)
        # 17 = wdExportFormatPDF
        doc_com.SaveAs(abs_out, FileFormat=17)
        doc_com.Close()
        print(f"PDF generado exitosamente en: {abs_out}")
    finally:
        word.Quit()

if __name__ == "__main__":
    fill_and_export()
